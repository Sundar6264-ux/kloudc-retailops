"""Download a purchase order as Excel, PDF or CSV.

The owner (Oct 22): "export generated po in multiple formats: Excel, pdf,
csv ... but the backend will store in whatever the format we are using right
now". So nothing new is stored: a PO stays what it is today (a placed-order
row's items_json, a Send PO log row's items_json, or the live cart lines)
and the file is made from it on request. Read-only - no row is written,
which is also why the demo login may download.

Every format has the same columns as the PDF sent to vendors - S.No,
Supplier Code, Item Name, Qty - in the same order (grouped into the same
picking sections, so S.No matches the PDF). When a line has the optional
case fields (`cases` / `case_size`), "Cases" and "Units per case" are added,
then "Ordered in" ("cases" / "units") so each line's Qty reads the same as
on the PDF ("2 cases" / "3 units").
"""

import csv
import datetime as dt
import io
import re
from typing import Optional

from sqlalchemy.orm import Session

from app.services.csv_util import neutralize_formula
from app.services.po_export import _format_qty, has_case_fields, line_in_cases

# (format, media type) - a tuple: module-level state stays read-only (tenancy lint).
FORMATS = (
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("pdf", "application/pdf"),
    ("csv", "text/csv; charset=utf-8"),
)
BAD_FORMAT = "Pick a file type: Excel (xlsx), PDF or CSV."
NO_LINES = "This order has no items to download."


class PoDownloadError(Exception):
    def __init__(self, detail: str, status_code: int = 400):
        super().__init__(detail)
        self.detail = detail
        self.status_code = status_code


def normalize_format(value: Optional[str]) -> str:
    fmt = (value or "pdf").strip().lower().lstrip(".")
    if fmt in ("excel", "xls"):
        fmt = "xlsx"
    if fmt not in dict(FORMATS):
        raise PoDownloadError(BAD_FORMAT)
    return fmt


def filename(supplier: str, day: dt.date, fmt: str) -> str:
    """PO-<Vendor>-<date>.<ext>: the vendor's name with anything but ASCII
    letters and digits turned into single dashes ("Raja Foods & Co." ->
    "Raja-Foods-Co")."""
    vendor = re.sub(r"[^A-Za-z0-9]+", "-", supplier or "").strip("-") or "Vendor"
    return f"PO-{vendor[:60].rstrip('-')}-{day.isoformat()}.{fmt}"


def _columns(with_cases: bool) -> list:
    cols = ["S.No", "Supplier Code", "Item Name", "Qty"]
    return cols + ["Cases", "Units per case", "Ordered in"] if with_cases else cols


def _number(value):
    """A quantity as a number for Excel (whole numbers without ".0"); text
    left as it is."""
    try:
        f = float(value)
    except (TypeError, ValueError):
        return "" if value is None else str(value)
    return int(f) if f.is_integer() else f


def _grouped(db: Session, supplier: str, lines: list) -> list:
    from app.services.po_send import group_po_lines

    return group_po_lines(db, supplier, lines) or [(None, lines)]


def _rows(groups: list, with_cases: bool) -> list:
    """[(section or None, [cells])] in printed order, S.No counting across
    sections like the PDF."""
    out = []
    idx = 0
    for section, section_lines in groups:
        for line in section_lines:
            idx += 1
            cells = [idx, line.get("supplier_item_id") or "", line.get("name") or "", line.get("qty")]
            if with_cases:
                cells += [line.get("cases"), line.get("case_size"), "cases" if line_in_cases(line) else "units"]
            out.append((section, cells))
    return out


def _csv(rows: list, with_cases: bool) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_columns(with_cases))
    for _section, cells in rows:
        # Item names and codes are typed text: a cell that a spreadsheet would
        # run as a formula gets a leading apostrophe (csv_util).
        writer.writerow(
            [cells[0], neutralize_formula(str(cells[1])), neutralize_formula(str(cells[2])), *(_format_qty(c) for c in cells[3:6]), *cells[6:]]
        )
    # BOM so Excel opens the file as UTF-8 (item names with accents).
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _as_text(ws, row: int) -> None:
    """openpyxl stores a string starting with "=" as a formula; item names
    and codes are typed text, so keep them text."""
    for cell in ws[row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _xlsx(supplier: str, day: dt.date, rows: list, with_cases: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from app.store import settings_view as sv

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase order"
    cols = _columns(with_cases)
    bold = Font(bold=True)
    ws.append(["Purchase order"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Supplier", supplier])
    _as_text(ws, ws.max_row)
    ws.append(["Prepared", day.isoformat()])
    store = sv.store_display_name()
    if store:
        ws.append(["From", store])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = bold
    ws.append([])
    ws.append(cols)
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="FBE4CC")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = bold
        cell.fill = fill
    current = object()
    section_fill = PatternFill("solid", fgColor="F3DCC2")
    for section, cells in rows:
        if section and section != current:
            ws.append([str(section).upper()])
            r = ws.max_row
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
            ws.cell(row=r, column=1).font = bold
            ws.cell(row=r, column=1).fill = section_fill
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        current = section
        ws.append([cells[0], cells[1], cells[2], *(_number(c) for c in cells[3:6]), *cells[6:]])
        _as_text(ws, ws.max_row)
    widths = [7, 18, 48, 10, 10, 15, 12]
    for i, width in enumerate(widths[: len(cols)]):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(db: Session, supplier: str, items: list, day: dt.date, fmt: str) -> tuple:
    """(bytes, media type, filename) for a PO's lines in `fmt`."""
    from app.services.po_send import build_po_pdf, po_lines

    fmt = normalize_format(fmt)
    items = [i for i in items if isinstance(i, dict) and str(i.get("name") or "").strip()]
    if not items:
        raise PoDownloadError(NO_LINES)
    for i in items:
        i.setdefault("qty", 0)
    if fmt == "pdf":
        content = build_po_pdf(db, supplier, items, prepared=day)
    else:
        lines = po_lines(items)
        with_cases = has_case_fields(lines)
        rows = _rows(_grouped(db, supplier, lines), with_cases)
        content = _csv(rows, with_cases) if fmt == "csv" else _xlsx(supplier, day, rows, with_cases)
    return content, dict(FORMATS)[fmt], filename(supplier, day, fmt)
