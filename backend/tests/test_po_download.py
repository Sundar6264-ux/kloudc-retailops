"""Download a PO as Excel / PDF / CSV (owner, Oct 22): the purchase order
log, the Send PO log and the cart, the same columns as the vendor PDF, plain
file names, the cart permission, the demo login allowed (read-only), and the
optional Cases / Units per case columns."""

import csv
import datetime as dt
import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.models import CartItem, PlacedOrder
from app.po_send_models import PoSend
from app.services import po_download
from app.timeutil import restaurant_today
from tests.conftest import client_for, make_user

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LINES = [{"name": "Toor Dal 4 lb", "supplier_item_id": "TD4", "qty": 2}, {"name": "Rooh Afza", "supplier_item_id": "RA2", "qty": 1.5}]
WHEN = dt.datetime(2026, 10, 1, 15, 0, tzinfo=dt.timezone.utc)


@pytest.fixture
def buyer(db):
    return client_for(db, make_user(db, "buyer", features=["purchase_order_cart"]))


def _order(db, lines=LINES, supplier="Raja Foods & Co."):
    row = PlacedOrder(
        supplier=supplier, item_count=len(lines), items_json=json.dumps(lines), placed_by_name="boss",
        placed_at=WHEN, status="placed",
    )
    db.add(row)
    db.commit()
    return row


def _send(db, lines=LINES, supplier="Raja Foods & Co."):
    row = PoSend(
        idempotency_key="k1", supplier=supplier, status="sent", item_count=len(lines), items_json=json.dumps(lines),
        email_to_json=json.dumps(["rep@vendor.test"]), created_at=WHEN,
    )
    db.add(row)
    db.commit()
    return row


def _pdf_text(content):
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)
    return " ".join(text.split())


def _xlsx_rows(content):
    ws = load_workbook(io.BytesIO(content)).active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _table(rows):
    """The rows from the column header on, without section headings."""
    start = next(i for i, r in enumerate(rows) if r and r[0] == "S.No")
    return [rows[start]] + [r for r in rows[start + 1 :] if r and isinstance(r[0], int)]


def test_order_as_xlsx(db, buyer):
    row = _order(db)
    r = buyer.get(f"/api/ordering/orders/{row.id}/export?format=xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == 'attachment; filename="PO-Raja-Foods-Co-2026-10-01.xlsx"'
    rows = _xlsx_rows(r.content)
    assert rows[0][0] == "Purchase order"
    assert ["Supplier", "Raja Foods & Co."] == list(rows[1][:2])
    table = _table(rows)
    assert list(table[0]) == ["S.No", "Supplier Code", "Item Name", "Qty"]
    body = sorted((t[1], t[2], t[3]) for t in table[1:])
    assert body == [("RA2", "Rooh Afza", 1.5), ("TD4", "Toor Dal 4 lb", 2)]
    assert sorted(t[0] for t in table[1:]) == [1, 2]


def test_order_as_pdf_has_lines_and_its_own_date(db, buyer):
    row = _order(db)
    r = buyer.get(f"/api/ordering/orders/{row.id}/export?format=pdf")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].endswith('filename="PO-Raja-Foods-Co-2026-10-01.pdf"')
    text = _pdf_text(r.content)
    for word in ("PURCHASE ORDER", "Raja Foods & Co.", "Toor Dal 4 lb", "TD4", "Rooh Afza", "1.5", "Prepared: 2026-10-01"):
        assert word in text, word
    assert "Cases" not in text and "Units per case" not in text


def test_order_as_csv(db, buyer):
    row = _order(db)
    r = buyer.get(f"/api/ordering/orders/{row.id}/export?format=csv")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"].endswith('filename="PO-Raja-Foods-Co-2026-10-01.csv"')
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["S.No", "Supplier Code", "Item Name", "Qty"]
    assert [r[0] for r in rows[1:]] == ["1", "2"]
    assert sorted(r[1:] for r in rows[1:]) == [["RA2", "Rooh Afza", "1.5"], ["TD4", "Toor Dal 4 lb", "2"]]


def test_default_is_pdf_and_bad_format_is_plain(db, buyer):
    row = _order(db)
    assert buyer.get(f"/api/ordering/orders/{row.id}/export").headers["content-type"] == "application/pdf"
    assert buyer.get(f"/api/ordering/orders/{row.id}/export?format=excel").headers["content-type"] == XLSX
    r = buyer.get(f"/api/ordering/orders/{row.id}/export?format=docx")
    assert r.status_code == 400 and r.json()["detail"] == po_download.BAD_FORMAT


def test_missing_and_empty_orders(db, buyer):
    assert buyer.get("/api/ordering/orders/999/export?format=csv").status_code == 404
    outside = _order(db, lines=[])
    r = buyer.get(f"/api/ordering/orders/{outside.id}/export?format=csv")
    assert r.status_code == 400 and r.json()["detail"] == po_download.NO_LINES


def test_sent_po_all_formats(db, buyer):
    row = _send(db)
    for fmt, ctype in (("xlsx", XLSX), ("pdf", "application/pdf"), ("csv", "text/csv")):
        r = buyer.get(f"/api/ordering/po-sends/{row.id}/export?format={fmt}")
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(ctype)
        assert r.headers["content-disposition"].endswith(f'filename="PO-Raja-Foods-Co-2026-10-01.{fmt}"')
        assert b"rep@vendor.test" not in r.content
    assert buyer.get("/api/ordering/po-sends/999/export").status_code == 404


def test_cart_po_matches_send_lines_and_is_not_logged(db, buyer):
    db.add(CartItem(supplier="Shaz", item_id="i1", name="Shan Masala", supplier_item_id="SH1", qty=3, case_of=1.0))
    db.add(CartItem(supplier="Shaz", item_id="i2", name="Zero line", supplier_item_id="Z", qty=0, case_of=1.0))
    db.commit()
    r = buyer.get("/api/ordering/cart/export", params={"supplier": "Shaz", "format": "xlsx"})
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].endswith(f'filename="PO-Shaz-{restaurant_today().isoformat()}.xlsx"')
    table = _table(_xlsx_rows(r.content))
    assert [list(t) for t in table[1:]] == [[1, "SH1", "Shan Masala", 3]]
    assert db.query(PlacedOrder).count() == 0, "a download is only a look"
    r = buyer.get("/api/ordering/cart/export", params={"supplier": "Nobody", "format": "csv"})
    assert r.status_code == 400 and "no items in the cart" in r.json()["detail"]
    r = buyer.get("/api/ordering/cart/export", params={"supplier": "Shaz", "format": "doc"})
    assert r.status_code == 400 and r.json()["detail"] == po_download.BAD_FORMAT


def test_case_columns_only_when_lines_have_them(db, buyer):
    lines = [
        {"name": "Basmati Rice 10 lb", "supplier_item_id": "BR10", "qty": 4, "cases": 2, "case_size": 2},
        {"name": "Ghee", "supplier_item_id": "G1", "qty": 1},
    ]
    row = _order(db, lines=lines, supplier="Acme")
    table = _table(_xlsx_rows(buyer.get(f"/api/ordering/orders/{row.id}/export?format=xlsx").content))
    assert list(table[0]) == ["S.No", "Supplier Code", "Item Name", "Qty", "Cases", "Units per case"]
    by_name = {t[2]: list(t) for t in table[1:]}
    assert by_name["Basmati Rice 10 lb"][3:] == [4, 2, 2]
    assert by_name["Ghee"][4:] in ([None, None], ["", ""])
    rows = list(csv.reader(io.StringIO(buyer.get(f"/api/ordering/orders/{row.id}/export?format=csv").content.decode("utf-8-sig"))))
    assert rows[0][-2:] == ["Cases", "Units per case"]
    text = _pdf_text(buyer.get(f"/api/ordering/orders/{row.id}/export?format=pdf").content)
    assert "Cases" in text and "Units per case" in text
    # case_of (the cart's own field) alone doesn't add the columns.
    plain = _order(db, lines=[{"name": "Ghee", "supplier_item_id": "G1", "qty": 1, "case_of": 12}], supplier="Acme")
    table = _table(_xlsx_rows(buyer.get(f"/api/ordering/orders/{plain.id}/export?format=xlsx").content))
    assert list(table[0]) == ["S.No", "Supplier Code", "Item Name", "Qty"]


def test_needs_the_cart_permission(db):
    row = _order(db)
    other = client_for(db, make_user(db, "projector", features=["supplier_projection"]))
    for path in (
        f"/api/ordering/orders/{row.id}/export",
        f"/api/ordering/po-sends/{_send(db).id}/export",
        "/api/ordering/cart/export?supplier=Raja",
    ):
        assert other.get(path).status_code == 403, path


def test_demo_login_can_download(db):
    from app.demo import is_demo_user

    demo = client_for(db, make_user(db, "app_store_review", features=["purchase_order_cart"]))
    assert is_demo_user(demo.user)
    row = _order(db)
    send = _send(db)
    assert demo.get(f"/api/ordering/orders/{row.id}/export?format=xlsx").status_code == 200
    assert demo.get(f"/api/ordering/po-sends/{send.id}/export?format=csv").status_code == 200


def test_filename_is_plain():
    day = dt.date(2026, 10, 22)
    assert po_download.filename("Raja Foods", day, "xlsx") == "PO-Raja-Foods-2026-10-22.xlsx"
    assert po_download.filename("  Café / Deli #2 ", day, "pdf") == "PO-Caf-Deli-2-2026-10-22.pdf"
    assert po_download.filename("", day, "csv") == "PO-Vendor-2026-10-22.csv"


def test_formula_looking_text_stays_text(db, buyer):
    row = _order(db, lines=[{"name": "=HYPERLINK(\"http://x\")", "supplier_item_id": "@SUM(A1)", "qty": 1}], supplier="Acme")
    rows = list(csv.reader(io.StringIO(buyer.get(f"/api/ordering/orders/{row.id}/export?format=csv").content.decode("utf-8-sig"))))
    assert rows[1][1:3] == ["'@SUM(A1)", "'=HYPERLINK(\"http://x\")"]
    ws = load_workbook(io.BytesIO(buyer.get(f"/api/ordering/orders/{row.id}/export?format=xlsx").content)).active
    cells = [c for r in ws.iter_rows() for c in r if c.value == '=HYPERLINK("http://x")']
    assert cells and all(c.data_type == "s" for c in cells)
